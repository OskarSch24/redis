#!/usr/bin/env python3
"""Konvertiert eine vom Job-Hunter generierte Excel-Datei in eine HTML-Vorschau.

Usage:
    python3 xlsx_to_html.py <input.xlsx> [output.html]

Wenn output.html nicht angegeben wird, landet sie neben der input.xlsx
mit gleichem Namen (nur Endung .html).

Zweck: Mac-Default kann .xlsx nicht öffnen wenn weder Numbers noch Excel
installiert ist. Diese HTML-Variante öffnet jeder Browser, URLs sind
klickbar, jedes Sheet kriegt einen Anker für die Sticky-Navigation.
"""

from __future__ import annotations

import html
import sys
from pathlib import Path

from openpyxl import load_workbook

_CSS = """
body{font-family:-apple-system,BlinkMacSystemFont,sans-serif;max-width:1400px;
     margin:20px auto;padding:0 16px;color:#222;}
h1{margin-top:8px;}
h2{margin-top:32px;border-bottom:2px solid #888;padding-bottom:4px;}
table{border-collapse:collapse;width:100%;font-size:13px;margin-bottom:24px;}
th{background:#2c3e50;color:#fff;padding:8px;text-align:left;position:sticky;top:48px;}
td{padding:6px 8px;border-bottom:1px solid #eee;vertical-align:top;}
tr:nth-child(even){background:#fafafa;}
tr:hover{background:#fff3cd;}
a{color:#0066cc;text-decoration:none;}
a:hover{text-decoration:underline;}
.nav{position:sticky;top:0;background:#fff;padding:8px 0;border-bottom:2px solid #ccc;
     z-index:10;}
.nav a{margin-right:8px;padding:4px 10px;background:#eee;border-radius:4px;
       font-weight:bold;font-size:12px;}
.meta{color:#666;font-size:12px;margin-bottom:16px;}
"""


def xlsx_to_html(xlsx_path: Path, html_path: Path | None = None) -> Path:
    """Konvertiert ein xlsx in eine HTML-Vorschau und gibt den Output-Pfad zurück."""
    if html_path is None:
        html_path = xlsx_path.with_suffix(".html")

    wb = load_workbook(xlsx_path, data_only=True)

    parts: list[str] = [
        "<!DOCTYPE html><html lang='de'><head>",
        "<meta charset='utf-8'>",
        f"<title>{html.escape(xlsx_path.name)}</title>",
        f"<style>{_CSS}</style>",
        "</head><body>",
        f"<h1>{html.escape(xlsx_path.stem)}</h1>",
        f"<div class='meta'>Quelle: {html.escape(str(xlsx_path))}</div>",
        "<div class='nav'>",
        " ".join(f"<a href='#{html.escape(s)}'>{html.escape(s)}</a>" for s in wb.sheetnames),
        "</div>",
    ]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        row_count = max(ws.max_row - 1, 0)
        parts.append(f"<h2 id='{html.escape(sheet_name)}'>")
        parts.append(f"{html.escape(sheet_name)} <span style='font-weight:400;color:#666;font-size:14px;'>")
        parts.append(f"({row_count} Zeilen)</span></h2>")
        parts.append("<table>")
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                parts.append("<tr>")
                for c in row:
                    parts.append(f"<th>{html.escape(str(c) if c is not None else '')}</th>")
                parts.append("</tr>")
                continue
            parts.append("<tr>")
            for c in row:
                if c is None:
                    parts.append("<td></td>")
                elif isinstance(c, str) and (c.startswith("http://") or c.startswith("https://")):
                    short = c if len(c) <= 60 else c[:57] + "…"
                    parts.append(f"<td><a href='{html.escape(c)}' target='_blank' rel='noopener'>"
                                 f"{html.escape(short)}</a></td>")
                else:
                    parts.append(f"<td>{html.escape(str(c))}</td>")
            parts.append("</tr>")
        parts.append("</table>")

    parts.append("</body></html>")
    html_path.write_text("".join(parts), encoding="utf-8")
    return html_path


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 xlsx_to_html.py <input.xlsx> [output.html]", file=sys.stderr)
        sys.exit(1)
    src = Path(sys.argv[1]).expanduser().resolve()
    if not src.exists():
        print(f"ERROR: Datei nicht gefunden: {src}", file=sys.stderr)
        sys.exit(1)
    dst = Path(sys.argv[2]).expanduser().resolve() if len(sys.argv) >= 3 else None
    out = xlsx_to_html(src, dst)
    print(out)


if __name__ == "__main__":
    main()
