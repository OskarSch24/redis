"""Excel exporter — creates a multi-sheet workbook from job postings."""

from __future__ import annotations

import logging
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from schema import JobPosting

logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.error("openpyxl not installed. Run: pip install openpyxl")


HEADER_FILL = "1F4E79"
HEADER_FONT_COLOR = "FFFFFF"
SCORE_COLORS = {
    "excellent": "00B050",  # green
    "good": "92D050",       # light green
    "neutral": "FFEB9C",    # yellow
    "poor": "FFC7CE",       # light red
    "reject": "FF0000",     # red
}

COLUMNS = [
    ("Score", 8),
    ("Firma", 25),
    ("Titel", 40),
    ("Standort", 20),
    ("Land", 6),
    ("Remote", 12),
    ("Arbeitszeit", 14),
    ("Stunden/Woche", 14),
    ("Gehalt Min", 12),
    ("Gehalt Max", 12),
    ("Währung", 9),
    ("Kaltakquise", 12),
    ("Telefon", 10),
    ("Sprachen", 15),
    ("Skills", 35),
    ("Verifiziert", 12),
    ("Aktiv", 8),
    ("Veröffentlicht", 14),
    ("Quelle", 18),
    ("URL", 50),
]


def export_to_excel(
    postings: list[JobPosting],
    output_path: str,
    config: dict[str, Any],
) -> str:
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl is required. Install with: pip install openpyxl")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    sorted_postings = sorted(postings, key=lambda p: p.score, reverse=True)
    top50 = sorted_postings[:50]

    _create_jobs_sheet(wb, "⭐ Top 50", top50)
    _create_jobs_sheet(wb, "📋 Alle Jobs", sorted_postings)
    _create_category_sheet(wb, "📁 Nach Kategorie", sorted_postings)
    _create_companies_sheet(wb, "🏢 Firmen-Übersicht", sorted_postings)
    _create_sources_sheet(wb, "📊 Quellen-Statistik", sorted_postings)
    _create_config_sheet(wb, "⚙️ Konfiguration", config)

    wb.save(output_path)
    logger.info("[export] Saved %d jobs to %s", len(postings), output_path)
    return str(output_path)


def _create_jobs_sheet(wb: "Workbook", sheet_name: str, postings: list[JobPosting]):
    ws = wb.create_sheet(sheet_name)
    col_names = [c[0] for c in COLUMNS]
    col_widths = [c[1] for c in COLUMNS]

    _write_header_row(ws, col_names)

    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, posting in enumerate(postings, 2):
        row_data = posting.to_excel_row()
        for col_idx, col_name in enumerate(col_names, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ""))
            cell.alignment = Alignment(wrap_text=False, vertical="center")

            if col_name == "URL" and row_data.get("URL"):
                cell.hyperlink = row_data["URL"]
                cell.style = "Hyperlink"

            if col_name == "Score":
                cell.number_format = "0.0"
                score = posting.score
                if score >= 70:
                    fill_color = SCORE_COLORS["excellent"]
                elif score >= 50:
                    fill_color = SCORE_COLORS["good"]
                elif score >= 30:
                    fill_color = SCORE_COLORS["neutral"]
                elif score >= 0:
                    fill_color = SCORE_COLORS["poor"]
                else:
                    fill_color = SCORE_COLORS["reject"]
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    score_col = get_column_letter(col_names.index("Score") + 1)
    last_row = len(postings) + 1
    if last_row > 2:
        ws.conditional_formatting.add(
            f"{score_col}2:{score_col}{last_row}",
            ColorScaleRule(
                start_type="num", start_value=-50, start_color="FF0000",
                mid_type="num", mid_value=0, mid_color="FFEB9C",
                end_type="num", end_value=100, end_color="00B050",
            ),
        )


def _create_category_sheet(wb: "Workbook", sheet_name: str, postings: list[JobPosting]):
    ws = wb.create_sheet(sheet_name)
    _write_header_row(ws, ["Kategorie", "Anzahl Jobs", "Ø Score", "Bester Score"])

    categories: dict[str, list[JobPosting]] = defaultdict(list)
    for p in postings:
        cat = _categorize(p.title)
        categories[cat].append(p)

    for col_idx, width in enumerate([30, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, (cat, jobs) in enumerate(
        sorted(categories.items(), key=lambda x: len(x[1]), reverse=True), 2
    ):
        scores = [j.score for j in jobs]
        ws.cell(row=row_idx, column=1, value=cat)
        ws.cell(row=row_idx, column=2, value=len(jobs))
        ws.cell(row=row_idx, column=3, value=round(sum(scores) / len(scores), 1))
        ws.cell(row=row_idx, column=4, value=round(max(scores), 1))

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _create_companies_sheet(wb: "Workbook", sheet_name: str, postings: list[JobPosting]):
    ws = wb.create_sheet(sheet_name)
    _write_header_row(ws, ["Firma", "Anzahl Stellen", "Ø Score", "Land", "Quellen"])

    for col_idx, width in enumerate([30, 14, 10, 8, 20], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    by_company: dict[str, list[JobPosting]] = defaultdict(list)
    for p in postings:
        by_company[p.company or "(Unbekannt)"].append(p)

    for row_idx, (company, jobs) in enumerate(
        sorted(by_company.items(), key=lambda x: len(x[1]), reverse=True), 2
    ):
        scores = [j.score for j in jobs]
        sources = ", ".join(sorted({j.source for j in jobs}))
        country = jobs[0].country if jobs else ""
        ws.cell(row=row_idx, column=1, value=company)
        ws.cell(row=row_idx, column=2, value=len(jobs))
        ws.cell(row=row_idx, column=3, value=round(sum(scores) / len(scores), 1))
        ws.cell(row=row_idx, column=4, value=country)
        ws.cell(row=row_idx, column=5, value=sources)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _create_sources_sheet(wb: "Workbook", sheet_name: str, postings: list[JobPosting]):
    ws = wb.create_sheet(sheet_name)
    _write_header_row(ws, ["Quelle", "Anzahl Jobs", "Aktiv verifiziert", "Inaktiv", "Ø Score"])

    for col_idx, width in enumerate([22, 12, 16, 10, 10], 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    by_source: dict[str, list[JobPosting]] = defaultdict(list)
    for p in postings:
        by_source[p.source].append(p)

    for row_idx, (source, jobs) in enumerate(
        sorted(by_source.items(), key=lambda x: len(x[1]), reverse=True), 2
    ):
        scores = [j.score for j in jobs]
        active = sum(1 for j in jobs if j.is_active is True)
        inactive = sum(1 for j in jobs if j.is_active is False)
        ws.cell(row=row_idx, column=1, value=source)
        ws.cell(row=row_idx, column=2, value=len(jobs))
        ws.cell(row=row_idx, column=3, value=active)
        ws.cell(row=row_idx, column=4, value=inactive)
        ws.cell(row=row_idx, column=5, value=round(sum(scores) / len(scores), 1) if scores else 0)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _create_config_sheet(wb: "Workbook", sheet_name: str, config: dict[str, Any]):
    import yaml
    ws = wb.create_sheet(sheet_name)
    ws.column_dimensions["A"].width = 80

    ws["A1"] = "Konfiguration (Snapshot zum Zeitpunkt des Runs)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Erstellt am: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    config_str = yaml.dump(config, allow_unicode=True, default_flow_style=False)
    for row_idx, line in enumerate(config_str.split("\n"), 4):
        ws.cell(row=row_idx, column=1, value=line)
        ws.cell(row=row_idx, column=1).font = Font(name="Courier New", size=9)


def _write_header_row(ws: Any, columns: list[str]):
    header_fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    header_font = Font(color=HEADER_FONT_COLOR, bold=True)

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _categorize(title: str) -> str:
    lower = title.lower()
    categories = [
        ("Customer Success", ["customer success", "kundenerfolg", "client success"]),
        ("Account Management", ["account manager", "account management", "key account"]),
        ("Customer Support", ["customer support", "customer service", "kundendienst", "helpdesk"]),
        ("Sales", ["sales", "vertrieb", "business development"]),
        ("Onboarding", ["onboarding", "implementation"]),
        ("Operations", ["operations", "ops"]),
        ("Partner Management", ["partner", "channel"]),
        ("Product", ["product manager", "product owner"]),
        ("Marketing", ["marketing", "growth"]),
        ("Engineering", ["engineer", "developer", "software"]),
    ]
    for category, keywords in categories:
        if any(kw in lower for kw in keywords):
            return category
    return "Sonstiges"
